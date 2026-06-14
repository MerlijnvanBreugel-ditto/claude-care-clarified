"""Generate a pragmatic Word doc with all website copy for team feedback."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document()

# Minimal styling
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(4)

h1_style = doc.styles['Heading 1']
h1_style.font.size = Pt(18)
h1_style.font.color.rgb = RGBColor(0x0D, 0x16, 0x4F)

h2_style = doc.styles['Heading 2']
h2_style.font.size = Pt(14)
h2_style.font.color.rgb = RGBColor(0x0D, 0x16, 0x4F)

h3_style = doc.styles['Heading 3']
h3_style.font.size = Pt(11)
h3_style.font.color.rgb = RGBColor(0x58, 0xB6, 0xD2)


def add_status_badge(doc, status_text, color):
    """Add a colored status line."""
    p = doc.add_paragraph()
    run = p.add_run(status_text)
    run.font.size = Pt(9)
    run.font.color.rgb = color
    run.bold = True
    p.paragraph_format.space_after = Pt(8)


def add_separator(doc):
    p = doc.add_paragraph()
    run = p.add_run('─' * 60)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    run.font.size = Pt(8)


def add_copy_block(doc, label, text):
    """Add a labeled copy block."""
    p = doc.add_paragraph()
    if label:
        run = p.add_run(label + ': ')
        run.bold = True
        run.font.size = Pt(10)
    run = p.add_run(text)
    run.font.size = Pt(10)


def add_quote(doc, text):
    """Add indented quote text."""
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


# ============================================================
# COVER
# ============================================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Ditto Care')
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x0D, 0x16, 0x4F)
run.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Website Copy — All Pages')
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x58, 0xB6, 0xD2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Draft for team feedback · March 2026')
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x7C, 0x82, 0x97)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(20)
run = p.add_run('Language: English (NL adaptation to follow)')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x7C, 0x82, 0x97)

doc.add_page_break()

# ============================================================
# PAGE 1: HOMEPAGE
# ============================================================
doc.add_heading('1. Homepage', level=1)
add_status_badge(doc, 'STATUS: First validation done by Merlijn', RGBColor(0x22, 0x8B, 0x22))
add_separator(doc)

doc.add_heading('Hero', level=2)
add_copy_block(doc, 'Headline', 'Clarity when it matters most')
add_copy_block(doc, 'Subheadline', 'Record every medical appointment. Get a clear summary in language you understand. Share it with the people who matter. Because care is something you carry together.')
add_copy_block(doc, 'CTA', 'Download Ditto [App Store] [Google Play]')
add_copy_block(doc, 'Secondary CTA', 'Are you a healthcare professional? →')
add_copy_block(doc, 'Trust line', '75,000+ users · 4.7★ rating · Promoted by hundreds of professionals')
add_copy_block(doc, 'Press logos', 'NOS · AD · Hart van Nederland · Quote · 3FM')
add_separator(doc)

doc.add_heading('The Problem', level=2)
add_copy_block(doc, 'Section label', 'Sound familiar?')

doc.add_heading('Card 1 — "The details fade"', level=3)
add_quote(doc, 'You walk out of the appointment and think: "Wait, what did they say?" By the time you get home, the details are already slipping. It happens to almost everyone. And the more serious the conversation, the harder it is to hold onto.')

doc.add_heading('Card 2 — It gets overwhelming', level=3)
add_quote(doc, "Some health journeys are straightforward. Many aren't. Suddenly there are specialists, letters, test results, scattered across systems, written in language that wasn't meant for you. The more important the moment, the harder it is to keep up.")

doc.add_heading('Card 3 — Everyone wants to know', level=3)
add_quote(doc, '"How did it go?" Your partner asks. Then your mother. Then your friend. Explaining what happened once is hard enough. Repeating it for everyone is exhausting. And somewhere in the retelling, you realize: you don\'t have a clear picture yourself.')
add_separator(doc)

doc.add_heading('Product Showcase', level=2)
add_copy_block(doc, 'Headline', 'This is Ditto')
add_copy_block(doc, 'Subheadline', 'Everything from your doctor\'s visit, captured, clarified, and shared with the people who care about you.')

features = [
    ('Record any appointment', 'Every word from your consultation, waiting for you when you\'re ready. Listen back, reread, and never lose an important detail again.'),
    ('A clear summary, instantly', 'What was discussed, your medications, and what happens next, so you walk away understanding, not guessing.'),
    ('Prepared for what\'s next', 'Ditto suggests questions for your next visit based on what was discussed. You show up ready, your doctor notices.'),
    ('Complex language, explained', 'Medical letters and specialist terms, translated into language that actually makes sense. Snap a photo, get clarity.'),
    ('Your Care Circle', 'Invite the people who matter. They see your summaries, follow your journey, and can actually help. Instead of asking "how did it go?" and getting a half-remembered answer.'),
    ('For patients and caregivers', 'Whether you\'re managing your own health or following someone you love, Ditto keeps everyone on the same page.'),
]
for title, desc in features:
    doc.add_heading(title, level=3)
    add_quote(doc, desc)
add_separator(doc)

doc.add_heading('Social Proof', level=2)
reviews = [
    ('Michel (Android)', 'Family', 'March 2026', 'Amazing app, so clear for loved ones and for summarizing difficult, lengthy conversations. A must-have for anyone with family or friends who need care. Calmly listen back and reread what was said.'),
    ('Els (iOS)', 'Understanding', 'January 2026', 'This app is absolutely brilliant. Super clear. Explains medical terms. Transcribes what you record. Or translates a medical letter from a photo. A real lifesaver and must-have!'),
    ('Jet', 'Pregnancy', 'March 2026', "During my pregnancy I used Ditto to keep track of all my appointments. But when our son was born, it became truly indispensable! Everything the doctor said, I could listen back and easily share with my partner so he was always up to date. A must-have for every (expecting) mum!"),
    ('Bibaboris', 'Professional use', 'October 2025', 'As someone who works in disability care, this is a wonderful invention. You record the conversation with the doctor and the app creates a concise summary. That makes handovers easy and consistent for everyone!'),
    ('Freek', 'Caregiver', 'June 2025', "I've been using Ditto for months to keep track of my father's medical journey. Every conversation with any healthcare provider, we record it, get a summary, and have a clear record. The app is incredibly easy to use and listens like an extra pair of ears (without the emotions). The summary afterward is accurate and easy to share."),
    ('Frees', 'Medical letters', 'June 2025', 'I use it for understanding letters from hospitals, which are primarily in Dutch.'),
]
for author, theme, date, text in reviews:
    p = doc.add_paragraph()
    run = p.add_run(f'★★★★★ · {theme} · {date}')
    run.bold = True
    run.font.size = Pt(9)
    add_quote(doc, f'"{text}"\n— {author}')
add_separator(doc)

doc.add_heading('USPs', level=2)
usps = [
    ('Everything in one place', 'Your GP, your specialist, your midwife. All in one app. Every appointment, every summary, every document. No matter where you get care.'),
    ('Sharing is caring', "Your loved ones can receive and share care updates safely, with end-to-end encryption, within Ditto. You stay in control of who's in your Care Circle."),
    ('Private by design', "Your data lives on your device. When Ditto processes a recording, it uses secure EU servers, then deletes it immediately. We can't read your summaries.\n[Learn more about how we protect your data →]"),
    ('Built for you', 'Patient portals are built for hospitals. Ditto is built for you, and the people around you.'),
]
for title, desc in usps:
    doc.add_heading(title, level=3)
    add_quote(doc, desc)
add_separator(doc)

doc.add_heading('Narrative Video', level=2)
add_copy_block(doc, 'Above video', 'The Story Behind the Summary')
add_copy_block(doc, 'Below video', 'Care. Clarified.')
add_separator(doc)

doc.add_heading('Closing CTA', level=2)
add_copy_block(doc, 'Headline', 'Record your first appointment tomorrow')
add_copy_block(doc, 'Subtext', "Download Ditto, create an account in under a minute, and bring it to your next visit. It's free. No trial, no catch.")
add_copy_block(doc, 'CTA', '[App Store badge] [Google Play badge] [QR code]')
add_copy_block(doc, 'Secondary', 'Questions? Talk to the team →')
add_separator(doc)

doc.add_heading('FAQ', level=2)
faqs_home = [
    ('Do I need my doctor\'s permission to record?', "In the Netherlands, patients have the right to record their own medical consultations. We recommend mentioning it. Most doctors are supportive, and many already recommend Ditto to their patients."),
    ('What if the summary gets something wrong?', "Ditto's summaries are generated by AI and are meant to help you remember and understand, not to replace medical advice. You always have the full recording to check against. If something looks off, listen back to the original conversation."),
    ('Can my doctor or insurance company see my recordings?', "No. Your data stays on your device. We can't access your recordings or summaries, and we never share data with insurers, employers, or anyone else. You decide who sees what, through your Care Circle. [Read more about our approach to privacy →]"),
    ('Does it work with any type of doctor?', 'Yes. GP, specialist, midwife, physiotherapist, psychologist. Ditto works with any in-person consultation. It captures the conversation regardless of which healthcare provider you\'re seeing.'),
]
for q, a in faqs_home:
    p = doc.add_paragraph()
    run = p.add_run(f'Q: {q}')
    run.bold = True
    run.font.size = Pt(10)
    add_quote(doc, a)

add_copy_block(doc, '', '[See all frequently asked questions →]')

doc.add_page_break()

# ============================================================
# PAGE 2: FOR PROFESSIONALS
# ============================================================
doc.add_heading('2. For Professionals', level=1)
add_status_badge(doc, 'STATUS: First validation done by Merlijn', RGBColor(0x22, 0x8B, 0x22))
add_separator(doc)

doc.add_heading('Hero', level=2)
add_copy_block(doc, 'Headline', "Care doesn't end when the consultation does.")
add_copy_block(doc, 'Subheadline', 'Ditto helps your patients understand what was discussed, remember what matters, and share it with their family, so they come back prepared, not confused.')
add_copy_block(doc, 'CTA', 'Talk to a fellow doctor →')
add_copy_block(doc, 'Trust line', '75,000+ patients use Ditto · Recommended by GPs, oncologists, and midwives across the Netherlands')
add_copy_block(doc, 'Privacy one-liner', "Patient data stays on their device. GDPR compliant. We can't access summaries.")
add_separator(doc)

doc.add_heading('Why It Matters', level=2)
add_copy_block(doc, 'Section label', 'The problem is well-documented')
stats = [
    ('40-80%', 'of medical information provided by clinicians is forgotten immediately.', 'Kessels, 2003; Watson & McKinstry, 2009'),
    ('~50%', 'of patients leave consultations with significant misunderstandings about their diagnosis or treatment.', 'Engel et al., 2017'),
    ('', 'Informed patients have measurably better treatment adherence and health outcomes.', 'Stacey et al., Cochrane Review'),
]
for pct, text, cite in stats:
    p = doc.add_paragraph()
    if pct:
        run = p.add_run(pct + ' ')
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x58, 0xB6, 0xD2)
    run = p.add_run(text)
    run.font.size = Pt(10)
    p2 = doc.add_paragraph()
    run = p2.add_run(f'— {cite}')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x7C, 0x82, 0x97)

add_copy_block(doc, 'Transition', "This isn't a patient problem. It's a communication problem. Ditto helps close the gap.")
add_separator(doc)

doc.add_heading('How Ditto Works', level=2)
add_copy_block(doc, 'Section label', 'How it works — nothing changes for you')
steps = [
    ('Step 1 — Your patient records the consultation', "The patient opens Ditto and taps record. Nothing changes for you. No devices to set up, no integrations to configure, no workflow disruption. The patient focuses on the conversation."),
    ('Step 2 — AI generates a plain-language summary', "After the visit, Ditto produces a structured summary: what was discussed, what the diagnosis means, what the treatment plan is, what happens next. In language the patient understands."),
    ('Step 3 — The patient\'s Care Circle stays informed', 'Patients share their summary with family members through Ditto\'s Care Circle, reducing the "what did the doctor say?" phone calls. Everyone who needs to know gets the same clear picture, directly from the source.'),
    ('Step 4 — Your next appointment starts better', 'The patient comes back having reviewed their summary, prepared questions, and aligned with their family. Less time repeating, more time progressing.'),
]
for title, desc in steps:
    doc.add_heading(title, level=3)
    add_quote(doc, desc)
add_separator(doc)

doc.add_heading('Built Together with Professionals', level=2)
add_copy_block(doc, 'Section label', 'Part of the medical community')
add_quote(doc, 'Ditto has doctors on the team. We collaborate with academic hospitals, present at healthcare conferences, and work with clinicians to make sure the product holds up where it matters.')

doc.add_heading('In-house medical expertise', level=3)
add_quote(doc, "Our team includes practicing doctors who review AI output, shape the product, and make sure Ditto meets the standards they'd expect for their own patients.")

doc.add_heading('Research collaborations', level=3)
add_quote(doc, 'We work with academic institutions to validate and improve patient communication through AI.\n• UMCG — University Medical Center Groningen: patient communication and AI-assisted health literacy\n• [TODO: Add other active collaborations]')

doc.add_heading('Conference & community presence', level=3)
add_quote(doc, '[TODO: List specific conferences, dates, and topics]')

doc.add_heading("What we're working on", level=3)
add_quote(doc, "We invest heavily in AI robustness, making sure summaries are accurate, complete, and safe. Clinical validation research is underway. We'll publish results when we have them, not before.")

add_copy_block(doc, 'CTA', 'Interested in a research collaboration? Talk to our team →')
add_separator(doc)

doc.add_heading('Partners', level=2)
add_copy_block(doc, 'Section label', 'Who we work with')
partners = [
    'Menzis — Health insurer · 2M+ members',
    'UMCG — Academic hospital · Research collaboration',
    'Juvoly — Healthcare software · Integration partner',
    'CZ — Health insurer · Zeeland pilot [TODO: Verify]',
]
for p_text in partners:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(p_text)
    run.font.size = Pt(10)
add_copy_block(doc, 'CTA', 'Work with us → [dedicated partnerships page]')
add_separator(doc)

doc.add_heading('Professional Testimonials', level=2)
add_copy_block(doc, 'Section label', 'What your colleagues say')
add_quote(doc, '"I was skeptical at first. Another app? But then patients started coming back differently. Better prepared, calmer, with written questions. It genuinely improved the quality of our follow-up conversations."\n— [Name], GP, [Location]')
add_quote(doc, '"For oncology patients especially, the amount of information is overwhelming. Knowing they can go home and re-read everything in plain language, it takes pressure off both of us."\n— [Name], Oncologist, [Location]')
add_copy_block(doc, 'Video testimonial', 'Watch Dr. [Name] explain how Ditto works in practice [Featured video]')

p = doc.add_paragraph()
run = p.add_run('Note: All testimonials must be from real, named professionals willing to be quoted.')
run.font.size = Pt(8)
run.italic = True
run.font.color.rgb = RGBColor(0x7C, 0x82, 0x97)
add_separator(doc)

doc.add_heading('Privacy & Data Handling', level=2)
add_copy_block(doc, 'Section label', 'How we handle patient data')
privacy_points = [
    "On-device storage. Patient data stays on the patient's device.",
    'EU servers only. When AI processing is needed, we use secure servers in the European Union.',
    'Immediate deletion. Audio and data are deleted from our servers after processing.',
    'GDPR compliant. Full compliance with European data protection regulations.',
    "No access. We cannot read, access, or share patient summaries. By design.",
]
for pp in privacy_points:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(pp)
    run.font.size = Pt(10)
add_copy_block(doc, '', '[Read our full approach to Trust & Privacy →]')
add_separator(doc)

doc.add_heading('What to Tell Your Patients', level=2)
add_quote(doc, 'Next time a patient seems overwhelmed, you can say: "There\'s a free app called Ditto that can help you remember what we discussed today. It records the conversation and gives you a clear summary to take home."')
add_separator(doc)

doc.add_heading('CTA', level=2)
add_copy_block(doc, 'Headline', 'Talk to a fellow doctor')
add_copy_block(doc, 'Subtext', "Have questions about how Ditto works in your practice? Speak to a member of our medical team. No sales pitch, just a conversation between colleagues.")
add_copy_block(doc, 'CTA button', 'Book a conversation → [Calendly link]')
add_copy_block(doc, 'Secondary', 'Work with us →')
add_separator(doc)

doc.add_heading('FAQ', level=2)
faqs_pro = [
    ('Does Ditto give medical advice?', "No. Ditto captures what was said during the consultation and presents it in clear language. It doesn't interpret, diagnose, or recommend. Patients always have access to the full recording, and the app is clear that it supports understanding, not medical decision-making."),
    ('How do you ensure the summaries are correct?', "We built an extensive AI engine with evaluation and safeguard frameworks. If we suspect any errors in a summary, we don't return output. Our in-house doctors review and validate AI quality, and we collaborate with academic institutions and hospitals to continuously improve. Patients also always have the full recording to listen back to."),
    ('Can I access my patient recordings?', "Recordings and summaries are only stored on the device of the patient. Some patients may ask you to cross-check the generated summary. If they want, they can share it with you directly, either by adding you to their Care Circle or sending it through another communication channel."),
    ('Does it work with all types of consultations?', 'Ditto works with any in-person consultation. GP, specialist, midwife, physiotherapist, psychologist. It captures the conversation regardless of specialty or setting.'),
    ('How is patient data handled?', "All data stays on the patient's device. When AI processing is needed, we use secure EU servers and delete the data immediately after. We're fully GDPR compliant and can't access patient recordings or summaries. [Read more about Trust & Privacy →]"),
]
for q, a in faqs_pro:
    p = doc.add_paragraph()
    run = p.add_run(f'Q: {q}')
    run.bold = True
    run.font.size = Pt(10)
    add_quote(doc, a)

add_copy_block(doc, '', '[See all frequently asked questions →]')

doc.add_page_break()

# ============================================================
# PAGE 3: TRUST & PRIVACY
# ============================================================
doc.add_heading('3. Trust & Privacy', level=1)
add_status_badge(doc, 'STATUS: First validation done by Merlijn', RGBColor(0x22, 0x8B, 0x22))
add_separator(doc)

doc.add_heading('Hero', level=2)
add_copy_block(doc, 'Headline', 'Your health data, your rules.')
add_copy_block(doc, 'Subheadline', "Healthcare is deeply personal, and so is the data that comes with it. Ditto is built with a privacy-by-design approach: we can never read or access what you do and store in the app. Here's exactly how it works.")
add_copy_block(doc, 'Trust badges', 'GDPR · EU Servers · [Any additional certifications]')
add_separator(doc)

doc.add_heading('How Recording Works', level=2)
add_copy_block(doc, 'Section label', 'What happens when you tap record')
rec_steps = [
    ('Step 1 — You record your consultation', "You open Ditto and tap record. The audio is captured on your phone. Your doctor doesn't need to install anything or change their workflow."),
    ('Step 2 — Transcription happens securely in the Netherlands', 'The audio is sent to NEN-certified Juvoly infrastructure in the Netherlands for transcription. This turns your conversation into text, securely, within Dutch borders.'),
    ('Step 3 — AI generates your summary', 'The transcribed text is processed on EU-based Azure servers to produce a structured, plain-language summary: what was discussed, what it means, and what happens next.'),
    ('Step 4 — Your data is deleted from our servers', "After processing, all audio and transcription data is deleted from our servers. We don't keep it. We can't replay it. It's gone."),
    ('Step 5 — Your summary arrives on your device', 'The summary lives on your phone. It never leaves your device unless you choose to share it, through your Care Circle or any other way you prefer.'),
]
for title, desc in rec_steps:
    doc.add_heading(title, level=3)
    add_quote(doc, desc)
add_separator(doc)

doc.add_heading('How We Make Summaries', level=2)
add_copy_block(doc, 'Section label', 'What goes into the AI behind your summary')
add_quote(doc, "Ditto doesn't just run your recording through a generic AI model. We built a dedicated AI engine designed specifically for medical conversations.")

summary_items = [
    ('Structured output', "Every summary follows a consistent structure: what was discussed, your diagnosis explained, medications, next steps, and open questions. This isn't freeform. It's a framework designed with doctors and patients."),
    ('Safeguards and evaluation', "We built extensive evaluation and safeguard frameworks around our AI. Every summary is checked for quality and consistency. If our AI is uncertain about something, it won't guess. It will tell you that it cannot reliably provide a summary rather than risk giving you something wrong."),
    ('Validated with doctors', 'Our in-house doctors review AI output continuously. We collaborate with academic institutions and hospitals to validate the quality and accuracy of our summaries.'),
    ('White paper', 'We\'ve documented our approach to AI quality and validation in detail. [Read our white paper on AI quality validation →]'),
]
for title, desc in summary_items:
    doc.add_heading(title, level=3)
    add_quote(doc, desc)
add_separator(doc)

doc.add_heading('Data Storage', level=2)
add_copy_block(doc, 'Section label', 'Where your data lives')
storage_items = [
    ('Everything stays on your device', 'Your summaries, recordings, documents, and Care Circle updates are stored locally on your phone. Not on a server. Not in a cloud we control. On your device.'),
    ('You control access', 'Your app is protected by biometric login or passcode. No one, not even us, can access what\'s inside.'),
    ('EU servers for processing only', "When you record a consultation or scan a document, the data is temporarily sent to secure servers in the European Union for processing. After processing, it's deleted. The only thing that remains is the summary on your device."),
    ('Your backup, your choice', "Ditto uses your device's built-in backup (iCloud or Google Drive). We recommend enabling encrypted backups so you can recover your information if you lose your phone."),
]
for title, desc in storage_items:
    doc.add_heading(title, level=3)
    add_quote(doc, desc)
add_separator(doc)

doc.add_heading('Care Circle Encryption', level=2)
add_copy_block(doc, 'Section label', 'How sharing stays private')
add_quote(doc, "When you share a summary with your Care Circle, it's protected by end-to-end encryption using a protocol similar to Signal, the same standard used by the most secure messaging apps in the world.")
add_quote(doc, "This means: only you and the people you've invited can read what's shared. Not us. Not anyone in between. The encryption happens on your device before the data ever leaves it.")
add_quote(doc, "You stay in full control of who's in your Care Circle and what they can see. You can add or remove people at any time.")
add_separator(doc)

doc.add_heading('AI Transparency', level=2)
add_copy_block(doc, 'Section label', 'How we use AI, and how we keep it safe')

doc.add_heading("What Ditto's AI does", level=3)
ai_does = [
    'Transcribes your medical conversation into text',
    'Generates a plain-language summary with structure: what was discussed, what it means, what\'s next',
    'Translates medical jargon into everyday language',
    'Explains medical documents you scan with your camera',
]
for item in ai_does:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(item)
    run.font.size = Pt(10)

doc.add_heading('How we make it safe', level=3)
ai_safe = [
    "We never train on your data. In fact, we don't even store it after processing.",
    'EU infrastructure only. Transcription happens in the Netherlands via NEN-certified Juvoly. Summaries are generated on EU-based Azure servers.',
    'Validation and governance. We work with patients and doctors to extensively test our AI before releasing new versions. [Read our white paper →]',
    'Built-in uncertainty. If our AI is uncertain about something, it tells you, rather than guessing. No output is better than wrong output.',
]
for item in ai_safe:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(item)
    run.font.size = Pt(10)

doc.add_heading("What Ditto's AI does NOT do", level=3)
add_quote(doc, "Ditto does not provide medical advice, recommend treatments, or make decisions about your care. It's a translator, not a doctor. It makes information understandable. Your healthcare team provides the care.")

doc.add_heading('EU AI Act', level=3)
add_quote(doc, 'Ditto is committed to responsible AI development in accordance with the EU AI Act. [Read more about our AI compliance →]')
add_separator(doc)

doc.add_heading('What We Will Never Do', level=2)
add_copy_block(doc, 'Section label', 'Our commitments to you')
nevers = [
    'We will never sell your data. Not to insurers. Not to advertisers. Not to anyone.',
    "We will never share data with your insurer. Even if your insurer partners with us, they cannot see your personal health information.",
    "We will never train AI on your recordings. Your conversations are processed and deleted. They don't become training data.",
    'We will never access your summaries. They live on your device. We have no technical ability to read them.',
]
for item in nevers:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(item)
    run.font.size = Pt(10)
    run.bold = True
add_separator(doc)

doc.add_heading('GDPR Compliance', level=2)
add_copy_block(doc, 'Section label', 'Your rights under EU law')
add_quote(doc, 'Ditto is fully compliant with the General Data Protection Regulation (GDPR). This means:')
gdpr_items = [
    "Right to access. You can request a copy of any data we hold about you, though almost everything is already on your device.",
    'Right to deletion. You can delete your account and all associated data at any time.',
    'Right to portability. Your data belongs to you. You can export it whenever you want.',
    'Data Processing Agreement. Healthcare professionals and organizations using Ditto can request a DPA for formal compliance documentation.',
]
for item in gdpr_items:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(item)
    run.font.size = Pt(10)
add_quote(doc, 'The only information we store centrally is fully anonymous technical logging and quality metrics (kept on encrypted Azure servers in Sweden) needed to ensure the app works properly and safely.')
add_separator(doc)

doc.add_heading("Iman's Video", level=2)
add_copy_block(doc, '', 'Watch: How Ditto protects your data')
add_quote(doc, 'Iman, [role at Ditto], explains our approach to privacy in under 3 minutes. [Video player]')
add_separator(doc)

doc.add_heading('Certifications & Partners', level=2)
add_copy_block(doc, 'Section label', 'Who trusts Ditto')
add_copy_block(doc, 'Badges', '[GDPR, any ISO, any additional]')
add_copy_block(doc, 'Trusted by', 'Menzis (health insurer, 2M+ members) · UMCG (academic hospital) · Hundreds of healthcare professionals')
add_separator(doc)

doc.add_heading('CTA', level=2)
add_copy_block(doc, 'Headline', 'Questions about privacy or security?')
add_copy_block(doc, 'Subtext', "Reach out to Yury, our CTO, who genuinely loves talking about keeping your data safe. Because clarity shouldn't come at the cost of privacy.")
add_copy_block(doc, 'CTA', 'Talk to Yury → [contact link]')
add_copy_block(doc, 'Secondary', 'Read our full privacy policy → | Vulnerability disclosure →')

doc.add_page_break()

# ============================================================
# PAGE 4: OUR STORY
# ============================================================
doc.add_heading('4. Our Story', level=1)
add_status_badge(doc, 'STATUS: First validation done by Merlijn', RGBColor(0x22, 0x8B, 0x22))
add_separator(doc)

doc.add_heading('Hero', level=2)
add_copy_block(doc, 'Headline', 'Our Story')
add_separator(doc)

doc.add_heading('The Mission', level=2)
add_quote(doc, "Across Europe, millions of patients struggle to make sense of medical information and coordinate their care. Consultations are short, health literacy levels are low, and families often support loved ones without the tools to understand diagnoses, treatment plans, or next steps. This lack of clarity leads to anxiety, preventable follow-up visits, and a feeling of being lost in your own healthcare.")
add_quote(doc, "The numbers confirm what most people already feel: research shows that 40 to 80 percent of medical information is forgotten immediately after a consultation. The more serious the news, the less sticks. And what patients do remember, they pass on to their partner, their children, their parents, each time a little different, a little less precise. By the time the information reaches the people who need it, it's no longer the doctor's story. It's a story about the doctor.")
add_quote(doc, "Healthcare systems are designed for the providers. Patient portals speak in medical jargon. Records are scattered across hospitals and specialists. And the people who care about you, your family, your friends, the ones who show up and ask what they can do, are systematically left out of the information loop. They want to help, but they don't have what they need to help well.")
add_quote(doc, "We thought this could and should be different. Not by fixing the system from the inside, but by building something new on the outside. Something that starts with the patient and the people around them. Something that makes healthcare understandable, shareable, and navigable. So we built Ditto.")
add_separator(doc)

doc.add_heading('The Founders', level=2)
add_copy_block(doc, 'Section label', 'The same pattern')

doc.add_heading('Tobias Polak', level=3)
add_quote(doc, "Tobias went with a friend to the oncologist. His friend was told he had cancer. In the car afterward, they started going over the conversation. What exactly had the doctor said about the prognosis? Was she hopeful or cautious? And that remark about the side effects, should they be worried? They couldn't agree. They'd been in the same room, heard the same words, and remembered two different versions.")

doc.add_heading('Bart Voorn', level=3)
add_quote(doc, "Bart was diagnosed with type 1 diabetes at thirty. Daily blood sugar checks, insulin injections, watching what you eat. And the people around you need to know what to do when your sugar drops too low. That you start shaking, get confused. That someone needs to act fast. Bart tried to explain all of this at home to the people closest to him. But what he passed on was already his version. Everyone understood it slightly differently. And with diabetes, it matters that the people around you know exactly what to do.")

doc.add_heading('Merlijn van Breugel', level=3)
add_quote(doc, "Merlijn sees it with his 87-year-old grandmother, who is being treated for lung cancer at the Antoni van Leeuwenhoek hospital. She always brings someone to the consultation room. Someone who can remember what the doctor says on her behalf. But even that person doesn't remember everything. And what they relay to the rest of the family at home is already their own version. Not the doctor's story. A story about the doctor.")

p = doc.add_paragraph()
run = p.add_run("Three people, the same pattern. They decided to do something about it.")
run.bold = True
run.font.size = Pt(11)
add_separator(doc)

doc.add_heading('What Ditto Is', level=2)
add_copy_block(doc, 'Section label', 'What we\'re building')
add_quote(doc, "Ditto is a patient navigation platform that uses AI to help people understand and manage their care. Patients can record consultations, receive accurate plain-language summaries, prepare for appointments, and share important information with their Care Circle of family and friends.")
add_quote(doc, "What started as a tool to record and summarize a doctor's visit has grown into something bigger: a place where your complete care journey lives, where complex medical language becomes understandable, and where the people around you can actually help instead of just asking how it went.")
add_quote(doc, "Clinicians are noticing too. Hundreds of healthcare professionals now recommend Ditto to their patients, because better-informed patients lead to better conversations and better outcomes for everyone.")
add_separator(doc)

doc.add_heading('What We Believe', level=2)
beliefs = [
    ('The patient unlocks everything.', "Healthcare systems are built for the system. We build for the person inside it. When you empower the patient, the family can help, the doctor gets a better partner, and outcomes improve. The unlock is never a better dashboard for the provider. It's a clearer picture for the person going through it."),
    ('Health is never solo.', "A diagnosis doesn't land on one person. It lands on a family, a friend group, a care circle. Every product decision we make accounts for the people around the patient, not as an add-on, but as core to how care actually works. That's why 39% of our summaries are shared."),
    ('Simplicity is the feature.', "In healthcare, complexity is the default. We strip it away relentlessly. If a feature makes the product harder to understand, it doesn't ship. If a screen can say it in fewer words, it should. Simplicity isn't a design preference. It's how we deliver clarity to people in stressful moments."),
    ('Earn trust, don\'t extract it.', "We handle data from people at their most vulnerable. We will never sell it, monetize it, or use it in ways they wouldn't expect. Privacy isn't a feature. It's the price of admission."),
]
for title, desc in beliefs:
    doc.add_heading(title, level=3)
    add_quote(doc, desc)
add_separator(doc)

doc.add_heading('Our Journey', level=2)
add_copy_block(doc, 'Section label', 'A timeline of milestones')
milestones = [
    ('2024', 'Founded in Rotterdam. Three founders, one mission.'),
    ('Early 2025', 'First version of Ditto launches. Days after launch, #1 healthcare app in the Netherlands.'),
    ('2025', 'Covered by 15+ media outlets: NOS, AD, Hart van Nederland, Quote, 3FM.'),
    ('2025', 'Menzis partnership. Ditto becomes available to 2 million insured members.'),
    ('2025', 'UMCG research collaboration begins.'),
    ('2026', '75,000+ users. Hundreds of healthcare professionals recommending Ditto.'),
    ('2026', 'Care Circle launches. Sharing your health journey with the people who matter.'),
]
for date, desc in milestones:
    p = doc.add_paragraph()
    run = p.add_run(date + '  ')
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x58, 0xB6, 0xD2)
    run = p.add_run(desc)
    run.font.size = Pt(10)

p = doc.add_paragraph()
run = p.add_run('[TODO: Verify exact dates and add any missing milestones]')
run.font.size = Pt(9)
run.italic = True
run.font.color.rgb = RGBColor(0x7C, 0x82, 0x97)
add_separator(doc)

doc.add_heading('The Team', level=2)
add_quote(doc, 'A young yet ambitious and experienced team, with little patience for the status quo. Driven by a shared mission, excited about building fast and purposefully.')
add_copy_block(doc, 'Team card format', '[Photo] · [Name] · [Role] · [One personal line]')
add_separator(doc)

doc.add_heading('Careers', level=2)
add_copy_block(doc, 'Headline', 'Want to help us build this?')
add_quote(doc, "Ditto is building something that doesn't exist yet. We're a small team in Rotterdam with a big mission and the traction to match. If that excites you, we'd love to hear from you.")
add_copy_block(doc, 'CTA', 'See open positions → [Homerun link]')
add_separator(doc)

doc.add_heading('Press', level=2)
add_copy_block(doc, 'Section label', 'Ditto in the media')
add_copy_block(doc, 'Press logos', 'NOS · AD · Hart van Nederland · Quote · 3FM · NPO Radio 5 · RTL')
add_copy_block(doc, 'Press contact', 'For press inquiries: press@ditto.care')

doc.add_page_break()

# ============================================================
# PAGE 5: HELP & CONTACT
# ============================================================
doc.add_heading('5. Help & Contact', level=1)
add_status_badge(doc, 'STATUS: Not yet iterated on — first draft for review', RGBColor(0xCC, 0x77, 0x00))
add_separator(doc)

doc.add_heading('Hero', level=2)
add_copy_block(doc, 'Headline', 'How can we help?')
add_copy_block(doc, 'Subheadline', 'Find answers below, or get in touch directly.')
add_separator(doc)

doc.add_heading('FAQ — For Patients', level=2)
patient_faqs = [
    ('What is Ditto?', "Ditto is a free app that helps you understand and manage your healthcare. Record your doctor's visit, get a clear summary in plain language, understand medical letters, and share your health information with the people who matter to you."),
    ('How do I record a consultation?', "Open Ditto, tap the record button before your appointment starts, and tap stop when it's finished. Within minutes, you'll have a plain-language summary of everything that was discussed."),
    ('Do I need my doctor\'s permission to record?', "In the Netherlands, patients have the legal right to record their own medical consultations. We recommend letting your doctor know — most are supportive, and many already recommend Ditto to their patients."),
    ('How do I share a summary with my family?', 'Open any summary and tap "Share." You can invite people to your Care Circle — your partner, children, parents, or anyone you choose. Each person sees what you decide to share with them.'),
    ('Can I share different information with different people?', "Yes. Your Care Circle lets you control who sees what. Your partner might get the full summary. Your child might get only the key updates. You're always in control."),
    ("What if I don't want to share anything?", "That's fine. Ditto works perfectly as a personal tool. Sharing is optional — it's there when you want it."),
    ('How much does it cost?', 'Ditto is free. No subscriptions, no trial periods, no hidden costs.'),
    ('Is Ditto available in my language?', "Ditto currently works best in Dutch and English. The app can process conversations in other languages — the summary will be generated in your preferred language."),
]
for q, a in patient_faqs:
    p = doc.add_paragraph()
    run = p.add_run(f'Q: {q}')
    run.bold = True
    run.font.size = Pt(10)
    add_quote(doc, a)
add_separator(doc)

doc.add_heading('FAQ — For Healthcare Professionals', level=2)
pro_faqs = [
    ('How does Ditto work in my practice?', "Patients use Ditto on their own phone. There's nothing to install, integrate, or configure on your end. Patients tap record at the start of a consultation and get a summary afterward. Your workflow doesn't change."),
    ('Is the AI summary accurate?', "Ditto's AI generates summaries that capture the key points of a consultation in plain language. We collaborate with medical professionals and academic institutions (including UMCG) to continuously validate and improve accuracy."),
    ('What about patient consent?', 'Patients control their own recordings. In the Netherlands, patients have the legal right to record consultations. Ditto encourages patients to inform their doctor, and we provide guidance on how to approach this conversation.'),
    ('Do I need to do anything differently?', 'No. Speak naturally. The AI is designed to work with real medical conversations — it handles medical terminology, multiple speakers, and varying audio quality.'),
    ('Can I recommend Ditto to my patients?', "Yes. Many healthcare professionals already do. If you'd like materials, a demo, or want to discuss implementation, talk to our medical team."),
    ("I'd like to try it myself or get a demo. How?", "Book a conversation with one of our medical team members — they're fellow doctors who can walk you through it.\n[Talk to a fellow doctor →]"),
]
for q, a in pro_faqs:
    p = doc.add_paragraph()
    run = p.add_run(f'Q: {q}')
    run.bold = True
    run.font.size = Pt(10)
    add_quote(doc, a)
add_separator(doc)

doc.add_heading('FAQ — For Partners & Organizations', level=2)
partner_faqs = [
    ('How can my organization work with Ditto?', "We offer several types of partnerships: promotional partnerships (recommend Ditto to your members), research collaborations (academic and clinical studies), and guided pilots (structured implementation with support). [Learn more on our Professionals page →]"),
    ('Can we get a Data Processing Agreement?', 'Yes. Contact us for formal DPA documentation for your compliance requirements.'),
    ('How does the Menzis partnership work?', 'Menzis members can access Ditto as part of their insurer\'s health literacy offering. [Contact us for details on similar partnerships →]'),
]
for q, a in partner_faqs:
    p = doc.add_paragraph()
    run = p.add_run(f'Q: {q}')
    run.bold = True
    run.font.size = Pt(10)
    add_quote(doc, a)
add_separator(doc)

doc.add_heading('FAQ — Privacy & Data', level=2)
priv_faqs = [
    ('Where is my data stored?', 'On your device. Summaries, documents, and Care Circle connections live on your phone, not on our servers.'),
    ('What happens to my audio recording?', "When you record a consultation, the audio is temporarily sent to secure EU servers for processing. After your summary is generated, the audio is deleted from our servers. We can't replay it."),
    ('Can Ditto read my summaries?', "No. We don't have technical access to your summaries. They're stored on your device, encrypted, and under your control."),
    ('Can I delete my data?', 'Yes. You can delete any summary, recording, or your entire account at any time. Deletion is immediate and permanent.'),
    ('Does my insurer see my data?', 'No. Even if your insurer partners with Ditto, they cannot access your personal health information. Your data is yours.'),
]
for q, a in priv_faqs:
    p = doc.add_paragraph()
    run = p.add_run(f'Q: {q}')
    run.bold = True
    run.font.size = Pt(10)
    add_quote(doc, a)
add_separator(doc)

doc.add_heading('Contact', level=2)
add_copy_block(doc, 'Intro', "Can't find what you're looking for? We're here to help.")
add_copy_block(doc, 'Form fields', 'Your name · Email · I am a... [Patient / Professional / Partner / Press / Other] · Subject · Message')
add_copy_block(doc, 'For professionals', 'Prefer a conversation? [Talk to a fellow doctor →]')
add_copy_block(doc, 'Direct contact', 'Email: info@ditto.care · Phone: +31 85 115 5421')
add_separator(doc)

doc.add_heading('Video Resources', level=2)
add_copy_block(doc, 'Section label', 'Watch & learn')
add_copy_block(doc, 'Video 1', 'How Ditto protects your data — Iman explains our approach to privacy in under 3 minutes.')
add_copy_block(doc, 'Video 2', 'How Ditto works — See the app in action, from recording to sharing.')
add_copy_block(doc, 'Video 3', '[Additional video — TBD by team]')

# Save
output_path = '/Users/merlijnvanbreugel/Documents/GitHub/Ditto/claude-care-clarified/brand/website/Ditto Care - Website Copy - All Pages.docx'
doc.save(output_path)
print(f'Saved to: {output_path}')
